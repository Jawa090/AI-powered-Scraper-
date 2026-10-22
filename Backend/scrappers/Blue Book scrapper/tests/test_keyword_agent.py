import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
import pytest

pytest.skip("Standalone CLI, Runner, and Store modules were removed during BaseScraper modularization; parser tests are maintained in test_parsers.py", allow_module_level=True)

from openpyxl import load_workbook

from bluebook.cli import build_parser, main
from bluebook.crawler import Runner, _is_transient
from bluebook.export import export_all
from bluebook.ingest import ingest_html
from bluebook.legacy import _split_phones
from bluebook.parsers import parse_contacts, parse_listing, parse_profile
from bluebook.store import Store
from bluebook.urls import canonical, kind_of, profile_root_url, same_search


FIX = Path(__file__).parent / "fixtures"
LISTING_DUP_URL = "https://www.thebluebook.com/search.html?region=12&class=3580"
PROFILE_RICH_URL = "https://www.thebluebook.com/iProView/318510/harbor-mechanical/subcontractors/"


class ParserEnhancementTests(unittest.TestCase):
    def test_duplicate_listing_cards_preserve_branch_data(self):
        result = parse_listing((FIX / "listing_duplicates.html").read_text(), LISTING_DUP_URL)
        self.assertEqual(result["result_rows"], 2)
        self.assertEqual(result["per_page"], 2)
        self.assertEqual(result["duplicate_rows"], 1)
        self.assertEqual(result["max_page"], 7)
        self.assertEqual(len(result["companies"]), 1)
        company = result["companies"][0]
        self.assertEqual(company["profile_url"], "https://www.thebluebook.com/iProView/777123/widget-roofing/subcontractors/")
        self.assertEqual(company["email"], "bids@widget-roofing.test")
        self.assertIn("estimating@widget-roofing.test", company["extra_emails"])
        self.assertEqual(len(company["locations"]), 2)
        all_phones = {company["phone"], *company["extra_phones"]}
        self.assertTrue({"(214) 555-0101", "(817) 555-0102", "(817) 555-0199"} <= all_phones)
        self.assertEqual({r["class_code"] for r in result["related_searches"]}, {"4660", "9999"})

    def test_rich_profile_email_address_codes_and_service_areas(self):
        data = parse_profile((FIX / "profile_rich.html").read_text(), PROFILE_RICH_URL)
        self.assertEqual(data["email"], "hello@harbor-mechanical.test")
        self.assertIn("bids@harbor-mechanical.test", data["extra_emails"])
        self.assertEqual((data["street"], data["city"], data["state"], data["zip"]),
                         ("51 Church St.", "Freeport", "NY", "11520"))
        self.assertIn("1680", data["trade_codes"])
        self.assertIn("Electrical Contractors", data["trades"])
        self.assertIn("Nassau County, New York", data["service_areas"])
        self.assertIn("MWBE", data["certifications"])
        self.assertEqual(data["established"], "1999")
        self.assertEqual(data["company_size"], "25-49 Employees")

    def test_contacts_preserve_alternate_rows_and_partial_surname(self):
        html = """<html><head><title>Widget Co - Locations and Key Contacts</title></head><body><main>
        <div id='keyContactSection'><h4>Owners, Principals & Senior Executives</h4>
          <div><button>Casey Example</button><div>President</div><a href='tel:+12125550101'>(212) 555-0101</a></div>
          <div><button>Casey Example</button><div>President</div><a href='tel:+12125550102'>(212) 555-0102</a></div>
          <div><button>Donna .</button><div>Estimator</div>
               <a href='tel:+12125550103'>(212) 555-0103</a>
               <a href='mailto:Donna@widget.test'>Email</a></div>
        </div></main></body></html>"""
        data = parse_contacts(html, "https://www.thebluebook.com/iProView/99/widget/general-contractors/locations-contacts/",
                              expected_name="Widget Co")
        casey = [c for c in data["contacts"] if c["name"] == "Casey Example"]
        self.assertEqual(len(casey), 2)
        donna = next(c for c in data["contacts"] if c["name"] == "Donna .")
        self.assertEqual(donna["email"], "Donna@widget.test")


class StoreMergeTests(unittest.TestCase):
    def test_conflicts_and_cross_profile_duplicate_candidates(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            root = "https://www.thebluebook.com/iProView/101/widget-roofing/subcontractors/"
            store.upsert_company({
                "company_id": "101", "name": "Widget Roofing LLC", "pretty_url": root,
                "phone": "(214) 555-0101", "phone_e164": "+12145550101", "phone_source": "tel_link",
                "website": "https://widget.test", "email": "hello@widget.test",
            }, enrichment="contacts")
            store.upsert_company({
                "company_id": "101", "name": "Widget Roofing Group",
                "pretty_url": root + "locations-contacts/",
                "phone": "(817) 555-0102", "phone_e164": "+18175550102",
                "website": "https://other-widget.test",
            }, enrichment="listing")
            store.upsert_company({
                "company_id": "202", "name": "Widget Roofing, LLC",
                "pretty_url": "https://www.thebluebook.com/iProView/202/widget-roofing/subcontractors/",
                "phone": "(214) 555-0101", "phone_e164": "+12145550101", "website": "https://www.widget.test",
            }, enrichment="listing")
            rows = {r["company_id"]: r for r in store.companies_flat()}
            self.assertEqual(rows["101"]["profile_url"], root)
            self.assertEqual(rows["101"]["phone"], "(214) 555-0101")
            self.assertIn("(817) 555-0102", rows["101"]["all_phones"])
            candidates = store.duplicate_candidates()
            self.assertEqual({r["company_id"] for r in candidates}, {"101", "202"})
            self.assertTrue(any(r["duplicate_of_company_id"] for r in candidates))
            store.close()

    def test_contact_and_location_merge_preserves_variants_without_fake_rows(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            loc = {"label": "Main", "address_raw": "100 Main St., Dallas, TX 75201", "street": "100 Main St.",
                   "city": "Dallas", "state": "TX", "zip": "75201", "phone": "(214) 555-0101",
                   "phone_e164": "+12145550101"}
            store.upsert_company({
                "company_id": "101", "name": "Widget", "locations": [loc],
                "contacts": [{"name": "Casey Example", "department": "Owners", "title": "President",
                              "phone": "(214) 555-0101", "phone_e164": "+12145550101"},
                             {"name": "Casey Example", "department": "Owners", "title": "President",
                              "phone": "(214) 555-0102", "phone_e164": "+12145550102"}],
            })
            loc2 = dict(loc, label="Duplicate label", phone="(214) 555-0109", phone_e164="+12145550109")
            store.upsert_company({"company_id": "101", "locations": [loc2]})
            self.assertEqual(store.stats()["locations_rows"], 1)
            self.assertEqual(store.stats()["contacts_rows"], 2)
            self.assertIn("(214) 555-0109", store.companies_flat()[0]["all_phones"])
            store.close()

    def test_search_totals_do_not_shrink_from_short_final_page(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            url = store.add_search(LISTING_DUP_URL, "Roofing")
            store.note_search_totals(url, 100, 25)
            store.note_search_totals(url, 100, 5)
            row = store.db.execute("SELECT per_page,pages_expected FROM searches WHERE url=?", (url,)).fetchone()
            self.assertEqual(row, (25, 4))
            store.note_search_totals(url, 100, 5, max_page=7)
            row = store.db.execute("SELECT pages_expected FROM searches WHERE url=?", (url,)).fetchone()
            self.assertEqual(row[0], 7)
            store.close()

    def test_listing_validation_and_true_no_results(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            with self.assertRaises(ValueError):
                ingest_html(store, "<html><body>Unexpected shell</body></html>", LISTING_DUP_URL)
            result = ingest_html(store, "<html><body><h1>0 Companies Found</h1></body></html>", LISTING_DUP_URL)
            self.assertTrue(result["no_results"])
            self.assertEqual(result["companies"], 0)
            # An existing DB name must not make an empty profile shell "valid".
            store.upsert_company({"company_id": "99", "name": "Known Widget"})
            with self.assertRaises(ValueError):
                ingest_html(store, "<html><body>empty shell</body></html>",
                            "https://www.thebluebook.com/iProView/99/widget/general-contractors/", kind="profile")
            store.close()


class DynamicAndRunnerTests(unittest.TestCase):
    def test_dynamic_related_expansion_is_bounded_and_filters_diversity_variants(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            search = store.add_search(LISTING_DUP_URL, "Roofing", metadata={
                "keyword": "roofing contractors", "keywords": ["roofing contractors"],
                "trade": "Roofing Contractors", "class_code": "3580", "expansion": "balanced",
                "dynamic_expansion": True, "dynamic_max_depth": 1, "dynamic_max_related": 5,
                "location_label": "Texas East", "priority": 5,
            }, priority=5)
            result = ingest_html(store, (FIX / "listing_duplicates.html").read_text(), search,
                                 category="Roofing", search_url=search)
            self.assertEqual(result["related_searches_queued"], 1)
            queued = [r[0] for r in store.db.execute("SELECT url FROM searches ORDER BY url")]
            self.assertTrue(any("class=4660" in u for u in queued))
            self.assertFalse(any("class=9999" in u for u in queued))
            store.close()

    def test_runner_helpers_without_browser(self):
        runner = Runner(None, log=lambda _msg: None)
        self.assertTrue(_is_transient(TimeoutError("slow")))
        self.assertTrue(_is_transient(RuntimeError("HTTP 503")))
        self.assertFalse(_is_transient(ValueError("bad page")))
        self.assertTrue(runner._markers_ok("<h1>0 Companies Found</h1>", "listing", LISTING_DUP_URL))
        with self.assertRaises(ValueError):
            runner._validate_final_url(
                {"kind": "profile", "company_id": "101"},
                "https://www.thebluebook.com/iProView/101",
                "https://www.thebluebook.com/iProView/202/other/subcontractors/")


class UrlRegressionTests(unittest.TestCase):
    def test_bare_host_class_padding_and_subpage_root(self):
        self.assertEqual(canonical("www.thebluebook.com/search.html?region=3&class=0830"),
                         "https://www.thebluebook.com/search.html?class=830&region=3")
        contacts = "https://www.thebluebook.com/iProView/101/widget/subcontractors/locations-contacts/"
        self.assertEqual(kind_of(contacts), "contacts")
        self.assertEqual(profile_root_url(contacts),
                         "https://www.thebluebook.com/iProView/101/widget/subcontractors/")
        tab = "https://www.thebluebook.com/iProView/101/widget/subcontractors/key-contacts/"
        self.assertEqual(kind_of(tab), "profile")
        self.assertTrue(same_search(
            "https://www.thebluebook.com/search.html?region=1&class=1680",
            "https://www.thebluebook.com/search-build/search.html?region=1&class=1680&page=3&sortfacets=az"))


class ExportTests(unittest.TestCase):
    def test_lead_excel_columns_and_source(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            search = store.add_search(LISTING_DUP_URL, "Roofing", metadata={
                "keyword": "roofing contractors", "keywords": ["roofing contractors"],
                "trade": "Roofing Contractors", "class_code": "3580", "priority": 5,
                "dynamic_expansion": False,
            }, priority=5)
            ingest_html(store, (FIX / "listing_duplicates.html").read_text(), search,
                        category="Roofing", search_url=search)
            out = Path(td) / "exports"
            result = export_all(store, out, formats=("csv", "xlsx", "jsonl"))
            self.assertEqual(result["counts"]["companies"], 1)
            wb = load_workbook(out / "bluebook_leads.xlsx")
            self.assertEqual(wb.sheetnames[0], "Leads")
            ws = wb["Leads"]
            headers = [c.value for c in ws[1]]
            self.assertEqual(headers[:11], [
                "Company Name", "Construction Trade", "Phone Number", "Email Address", "Website URL",
                "Blue Book Profile URL", "Address", "City", "State", "ZIP Code", "Source URL"])
            row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
            self.assertEqual(row["Construction Trade"], "Roofing Contractors")
            self.assertEqual(row["Email Address"], "bids@widget-roofing.test")
            self.assertEqual(row["Source URL"], search)
            self.assertEqual(row["ZIP Code"], "75201")
            store.close()


class CliTests(unittest.TestCase):
    def test_db_may_appear_before_or_after_subcommand(self):
        parser = build_parser()
        before = parser.parse_args(["--db", "before.sqlite3", "plan", "--keyword", "roofing"])
        after = parser.parse_args(["plan", "--keyword", "roofing", "--db", "after.sqlite3"])
        self.assertEqual(before.db, "before.sqlite3")
        self.assertEqual(after.db, "after.sqlite3")

    def test_plan_json_and_scrape_dry_run_create_no_db(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = main(["plan", "--keyword", "roofing contractors", "--city", "Houston, TX",
                       "--expansion", "off", "--format", "json"])
        self.assertEqual(rc, 0)
        plan = json.loads(buf.getvalue())
        self.assertGreater(plan["count"], 0)
        with tempfile.TemporaryDirectory() as td:
            db = Path(td) / "dry.sqlite3"
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["scrape", "--db", str(db), "--keyword", "demolition", "--state", "TX",
                           "--expansion", "off", "--dry-run"])
            self.assertEqual(rc, 0)
            self.assertFalse(db.exists())

    def test_config_and_invalid_search_parameter(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / "config.json"
            cfg.write_text(json.dumps({"keywords": ["electrical contractors"], "states": ["NY"],
                                       "expansion": "off"}), encoding="utf-8")
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["plan", "--config", str(cfg), "--format", "json"])
            self.assertEqual(rc, 0)
            self.assertEqual(json.loads(buf.getvalue())["searches"][0]["class_code"], "1680")
        with contextlib.redirect_stderr(io.StringIO()):
            rc = main(["plan", "--keyword", "roofing", "--search-param", "class=2200"])
        self.assertEqual(rc, 2)


class LegacyPhoneTests(unittest.TestCase):
    def test_comma_and_dot_alternates_are_preserved(self):
        nums = _split_phones("(212) 555-0101, (212) 555-0102 · (212) 555-0103")
        self.assertEqual([n["e164"] for n in nums], ["+12125550101", "+12125550102", "+12125550103"])


if __name__ == "__main__":
    unittest.main()
