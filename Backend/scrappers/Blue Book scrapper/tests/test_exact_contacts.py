import tempfile
import unittest
from pathlib import Path
import pytest

pytest.skip("Standalone CLI, Store, and Excel export modules were removed during BaseScraper modularization; parser tests are maintained in test_parsers.py", allow_module_level=True)

from openpyxl import load_workbook

from bluebook.cli import main
from bluebook.export import (EXACT_CONTACT_COLUMNS, exact_contact_rows,
                             export_all, export_exact_contacts,
                             validate_exact_contact_rows)
from bluebook.names import normalize_person_name
from bluebook.parsers import parse_contacts, parse_profile
from bluebook.store import Store


FIX = Path(__file__).parent / "fixtures"
PROFILE_PEOPLE_URL = "https://www.thebluebook.com/iProView/424242/widget-builders/general-contractors/"


def _seed_exact_store(store: Store):
    store.upsert_company({
        "company_id": "101", "name": "Widget Builders Group",
        "pretty_url": "https://www.thebluebook.com/iProView/101/widget-builders/general-contractors/",
        "website": "https://widget-builders.test",
        "phone": "(214) 555-0100", "phone_e164": "+12145550100", "phone_source": "tel_link",
        "email": "info@widget-builders.test", "email_source": "json_ld",
        "locations": [{
            "label": "Widget Builders Group", "address_raw": "100 Main St., Dallas, TX 75201",
            "street": "100 Main St.", "city": "Dallas", "state": "TX", "zip": "75201",
            "phone": "(214) 555-0100", "phone_e164": "+12145550100", "is_main": True,
            "contact_name": "Robin Johnson", "contact_title": "Principal Contact",
        }],
        "contacts": [
            {"name": "Casey Example", "department": "Owners", "title": "President",
             "phone": "(214) 555-0100", "phone_e164": "+12145550100", "phone_source": "tel_link", "email": ""},
            {"name": "Casey Example", "department": "Estimating", "title": "Estimator",
             "phone": "(214) 555-0199", "phone_e164": "+12145550199", "phone_source": "tel_link",
             "email": "Casey@widget-builders.test"},
            {"name": "CASEY Example", "department": "Operations", "title": "Manager",
             "phone": "", "phone_e164": "", "email": "casey@widget-builders.test"},
            {"name": "Donna Smith Jr.", "department": "Estimating", "title": "Estimator",
             "phone": "", "phone_e164": "", "email": ""},
            {"name": "Donna Smith", "department": "Operations", "title": "Manager",
             "phone": "(214) 555-0198", "phone_e164": "+12145550198", "email": ""},
            {"name": "Alex Roe", "department": "Project Management", "title": "Project Manager",
             "phone": "", "phone_e164": "", "email": ""},
        ],
    }, enrichment="contacts")
    store.upsert_company({
        "company_id": "202", "name": "No People Builders",
        "pretty_url": "https://www.thebluebook.com/iProView/202/no-people/general-contractors/",
        "website": "https://no-people.test", "phone": "(310) 555-0101",
        "phone_e164": "+13105550101", "email": "hello@no-people.test", "email_source": "json_ld",
    }, enrichment="profile")
    # Company 303 tests assignment safety: the flat company phone/email are
    # actually Dana's sole direct values, so they must not fall back to Sam.
    store.upsert_company({
        "company_id": "303", "name": "Direct Line Builders",
        "pretty_url": "https://www.thebluebook.com/iProView/303/direct-line/general-contractors/",
        "phone": "(512) 555-0199", "phone_e164": "+15125550199",
        "email": "dana@direct-line.test", "email_source": "visible_text",
        "contacts": [
            {"name": "Dana Lee", "department": "Owners", "title": "Owner",
             "phone": "(512) 555-0199", "phone_e164": "+15125550199", "email": "dana@direct-line.test"},
            {"name": "Sam Lee", "department": "Estimating", "title": "Estimator",
             "phone": "", "phone_e164": "", "email": ""},
        ],
    }, enrichment="contacts")


class PeopleParserTests(unittest.TestCase):
    def test_profile_management_estimating_and_labelled_people(self):
        data = parse_profile((FIX / "profile_people.html").read_text(), PROFILE_PEOPLE_URL)
        by_name = {c["name"]: c for c in data["contacts"]}
        self.assertTrue({"Casey Example", "Donna Lee", "Robin Johnson", "Avery Stone"} <= set(by_name))
        self.assertEqual(by_name["Casey Example"]["title"], "President")
        self.assertEqual(by_name["Casey Example"]["email"], "Casey@widget-builders.test")
        self.assertEqual(by_name["Donna Lee"]["phone"], "(214) 555-0198")
        self.assertEqual(by_name["Avery Stone"]["department"], "Project Management")

    def test_contacts_all_sections_and_location_person(self):
        html = """<html><head><title>Widget Co - Locations and Key Contacts</title></head><body><main>
        <section><h2>Locations</h2><div id='locationsPage-1'><div>
          <h3>Widget Co</h3><div>100 Main St.<br>Dallas, TX 75201</div>
          <a href='tel:+12145550100'>(214) 555-0100</a>
          <p><strong>Location Person</strong><br>Branch Manager</p>
        </div></div></section>
        <section><h2>Key Contacts</h2><div id='keyContactSection'>
          <div><button>Key Person</button><div>President</div><a href='tel:+12145550101'>(214) 555-0101</a></div>
        </div></section>
        <section><h2>Company Management</h2>
          <div><strong>Management Person</strong><div>Operations Manager</div><a href='tel:+12145550102'>(214) 555-0102</a></div>
        </section>
        <section><h2>Principal Contacts</h2><div>Principal Contact: Principal Person</div></section>
        </main></body></html>"""
        data = parse_contacts(
            html, "https://www.thebluebook.com/iProView/99/widget/general-contractors/locations-contacts/",
            expected_name="Widget Co")
        names = {c["name"] for c in data["contacts"]}
        self.assertTrue({"Location Person", "Key Person", "Management Person", "Principal Person"} <= names)

    def test_person_name_normalization_is_conservative(self):
        self.assertEqual(normalize_person_name("Dr. Casey Example"), "casey example")
        self.assertNotEqual(normalize_person_name("Donna Smith"), normalize_person_name("Donna Smith Jr."))
        self.assertNotEqual(normalize_person_name("Anna Lee"), normalize_person_name("Ann A. Lee"))


class ExactExportTests(unittest.TestCase):
    def test_one_row_per_person_phone_email_and_blank_rules(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            _seed_exact_store(store)
            rows, stats = exact_contact_rows(store.companies_flat(), store.rows("contacts"), store.rows("locations"))
            validation = validate_exact_contact_rows(rows)
            self.assertTrue(validation["valid"], validation)
            self.assertEqual([list(r.keys()) for r in rows], [EXACT_CONTACT_COLUMNS] * len(rows))
            widget = {r["Person Name"]: r for r in rows if r["Company Name"] == "Widget Builders Group"}
            self.assertEqual(set(widget), {"Casey Example", "Donna Smith Jr.", "Donna Smith", "Alex Roe", "Robin Johnson"})
            self.assertEqual(widget["Casey Example"]["Phone Number"], "(214) 555-0199")
            self.assertEqual(widget["Casey Example"]["Email"].casefold(), "casey@widget-builders.test")
            self.assertEqual(widget["Donna Smith Jr."]["Phone Number"], "(214) 555-0100")
            self.assertEqual(widget["Donna Smith"]["Phone Number"], "(214) 555-0198")
            self.assertEqual(widget["Alex Roe"]["Email"], "info@widget-builders.test")

            no_people = [r for r in rows if r["Company Name"] == "No People Builders"]
            self.assertEqual(len(no_people), 1)
            self.assertEqual(no_people[0]["Person Name"], "")
            self.assertEqual(no_people[0]["Phone Number"], "(310) 555-0101")
            self.assertEqual(no_people[0]["Email"], "hello@no-people.test")

            direct = {r["Person Name"]: r for r in rows if r["Company Name"] == "Direct Line Builders"}
            self.assertEqual(direct["Sam Lee"]["Phone Number"], "")
            self.assertEqual(direct["Sam Lee"]["Email"], "")

            self.assertEqual(stats["unique_companies_processed"], 3)
            self.assertEqual(stats["individual_contacts_collected"], 7)
            # Store cross-section merging has already collapsed one no-conflict
            # Casey repeat; one conflicting direct-line variant remains for the
            # final one-person-per-row consolidation.
            self.assertEqual(stats["duplicate_person_records_removed"], 1)
            self.assertEqual(stats["duplicate_output_rows_removed"], 0)
            store.close()

    def test_exact_workbook_has_one_sheet_and_six_columns(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            _seed_exact_store(store)
            out = Path(td) / "exact.xlsx"
            result = export_exact_contacts(store, out)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Contacts"])
            ws = wb["Contacts"]
            self.assertEqual(ws.max_column, 6)
            self.assertEqual([c.value for c in ws[1]], EXACT_CONTACT_COLUMNS)
            self.assertEqual(ws.max_row, result["counts"]["output_rows"] + 1)
            self.assertTrue(result["counts"]["validation"]["valid"])
            store.close()

    def test_full_export_and_cli_include_exact_contacts(self):
        with tempfile.TemporaryDirectory() as td:
            db = Path(td) / "t.sqlite3"
            store = Store(db)
            _seed_exact_store(store)
            full = export_all(store, Path(td) / "exports", formats=("xlsx",))
            self.assertIn(str(Path(td) / "exports" / "bluebook_contacts_exact.xlsx"), full["files"])
            self.assertEqual(full["counts"]["individual_contacts"], 7)
            store.close()
            out = Path(td) / "cli_exact.xlsx"
            rc = main(["--db", str(db), "export-contacts", "--out", str(out)])
            self.assertEqual(rc, 0)
            self.assertTrue(out.exists())


class StorePersonMergeTests(unittest.TestCase):
    def test_cross_section_repeat_updates_blanks_without_duplicate(self):
        with tempfile.TemporaryDirectory() as td:
            store = Store(Path(td) / "t.sqlite3")
            store.upsert_company({"company_id": "1", "name": "Widget", "contacts": [
                {"name": "Casey Example", "department": "Owners", "title": "President",
                 "phone": "(212) 555-0101", "phone_e164": "+12125550101", "email": ""}]})
            store.upsert_company({"company_id": "1", "contacts": [
                {"name": "Casey Example", "department": "Estimating", "title": "Estimator",
                 "phone": "(212) 555-0101", "phone_e164": "+12125550101", "email": "Casey@widget.test"}]})
            rows = store.rows("contacts")
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["email"], "Casey@widget.test")
            stats = store.stats(include_duplicates=False)
            self.assertEqual(stats["individual_contacts"], 1)
            self.assertEqual(stats["raw_person_records"], 1)
            store.close()


if __name__ == "__main__":
    unittest.main()
